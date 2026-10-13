# -*- coding: utf-8 -*-
"""Gera o .docx do relatório a partir das seções em Markdown (relatorio/etapa1/*.md).

Uso:
    python relatorio/build.py --template "caminho/instruções_de_elaboração.docx" --out relatorio/build/relatorio-etapa1.docx

O template do professor não fica no repositório; cada integrante usa a sua cópia.
Os autores vêm de relatorio/autores.local.md (fora do Git; veja autores.exemplo.md).
Requer: pip install python-docx
"""
import argparse, copy, math, os, re, sys, glob

import docx
from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# Títulos do template que ficam só com o título nesta etapa
TITLE_ONLY = {
    "Etapa 2",
    "Aplicação, Testes e Validação de um Projeto Orientado a Testes com Programação Mobile",
    "Elaboração da Aplicação Mobile", "Elaboração dos Testes", "Validação do Projeto",
    "Ciclo de Desenvolvimento e Validação com TDD", "Vídeo de Apresentação",
}


# ------------------------------------------------------------------ leitura do Markdown
def parse_sections(folder):
    """Retorna {titulo_do_template: [blocos]} lendo todos os .md da pasta, em ordem."""
    sections, cur = {}, None
    for path in sorted(glob.glob(os.path.join(folder, "*.md"))):
        text = open(path, encoding="utf-8").read()
        text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
        buf = []

        def flush():
            nonlocal buf
            if buf and cur is not None:
                sections[cur].append(classify(buf, folder))
            buf = []

        for line in text.splitlines():
            if line.startswith("## "):
                flush()
                cur = line[3:].strip()
                sections.setdefault(cur, [])
            elif not line.strip():
                flush()
            else:
                buf.append(line.rstrip())
        flush()
    return sections


def classify(lines, folder):
    if all(l.startswith("- ") for l in lines):
        return ("ul", [l[2:].strip() for l in lines])
    if lines[0].startswith("|"):
        rows = []
        for l in lines:
            cells = [c.strip() for c in l.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                continue
            rows.append(cells)
        return ("table", rows)
    if lines[0].startswith("> "):
        return ("cap", " ".join(l[2:].strip() for l in lines))
    m = re.fullmatch(r"!\[.*?\]\((.+?)\)", lines[0])
    if m:
        return ("img", os.path.join(folder, m.group(1)))
    return ("p", " ".join(l.strip() for l in lines))


# ------------------------------------------------------------------ construção do docx
def build(template, out, sections, autores):
    doc = docx.Document(template)
    paras = list(doc.paragraphs)
    BODY = copy.deepcopy(paras[46]._p)
    BULLET = copy.deepcopy(paras[61]._p)

    def clear_runs(p):
        for child in list(p):
            if child.tag != W + "pPr":
                p.remove(child)

    def fill(p_el, text):
        para = docx.text.paragraph.Paragraph(p_el, doc._body)
        clear_runs(p_el)
        for chunk in re.split(r"(\[\[.*?\]\]|\*\*.*?\*\*)", text):
            if not chunk:
                continue
            if chunk.startswith("[["):
                r = para.add_run("[" + chunk[2:-2] + "]")
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            elif chunk.startswith("**"):
                r = para.add_run(chunk[2:-2]); r.bold = True
            else:
                para.add_run(chunk)
        return p_el

    def body(text):
        return fill(copy.deepcopy(BODY), text)

    def bullet(text):
        return fill(copy.deepcopy(BULLET), text)

    def caption(text):
        el = fill(copy.deepcopy(BODY), text)
        para = docx.text.paragraph.Paragraph(el, doc._body)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in para.runs:
            r.font.size = Pt(10)
        return el

    def picture(path):
        el = copy.deepcopy(BODY)
        clear_runs(el)
        para = docx.text.paragraph.Paragraph(el, doc._body)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # cabe na página: no máximo 15,5 cm de largura e 15 cm de altura
        from struct import unpack
        with open(path, "rb") as fh:
            head = fh.read(24)
        w_px, h_px = unpack(">II", head[16:24])  # PNG
        width_cm = min(15.5, 15.0 * w_px / h_px)
        para.add_run().add_picture(path, width=Cm(width_cm))
        return el

    def table(rows):
        ncols = len(rows[0])
        weights = []
        for ci in range(ncols):
            longest = max(len(re.sub(r"\[\[|\]\]|\*\*", "", r[ci])) for r in rows)
            weights.append(min(max(math.sqrt(longest), 3.0), 9.0))
        total = sum(weights)
        widths = [16.0 * w / total for w in weights]
        t = doc.add_table(rows=len(rows), cols=ncols)
        tbl = t._tbl
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement("w:" + edge)
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "808080")
            borders.append(e)
        tbl.tblPr.append(borders)
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                cell = t.cell(ri, ci)
                cell.text = ""
                p = cell.paragraphs[0]
                for chunk in re.split(r"(\[\[.*?\]\])", val):
                    if not chunk:
                        continue
                    if chunk.startswith("[["):
                        r = p.add_run("[" + chunk[2:-2] + "]"); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    else:
                        r = p.add_run(chunk)
                    r.font.size = Pt(9)
                    if ri == 0:
                        r.bold = True
                if ri == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "D9EAD3")
                    cell._tc.get_or_add_tcPr().append(shd)
                cell.width = Cm(widths[ci])
        tbl.getparent().remove(tbl)
        return tbl

    def elements(blocks):
        out_els = []
        for kind, val in blocks:
            if kind == "p":
                out_els.append(body(val))
            elif kind == "ul":
                out_els += [bullet(t) for t in val]
            elif kind == "table":
                out_els.append(table(val))
            elif kind == "cap":
                out_els.append(caption(val))
            elif kind == "img":
                out_els.append(picture(val))
        return out_els

    def text_of(key):
        return " ".join(v for k, v in sections[key] if k == "p")

    # --- cabeçalho
    runs = paras[0].runs
    runs[1].text = text_of("Título"); runs[0].text = ""; runs[2].text = ""
    for r in runs:
        r.font.color.rgb = RGBColor(0, 0, 0)

    names = [a[0] for a in autores]
    rs = paras[2].runs
    rs[0].text = names[0]
    if len(names) > 1: rs[2].text = ", " + names[1]
    if len(names) > 2: rs[4].text = ", " + names[2]
    rs[5].text = "1"
    for extra in names[3:]:
        r = paras[2].add_run(", " + extra); r.bold = True; r.font.size = rs[0].font.size
        s = paras[2].add_run("1"); s.bold = True; s.font.size = rs[0].font.size; s.font.superscript = True

    for r in paras[3].runs:
        if "Rua" in r.text:
            r.text = r.text.rstrip() + " Uruguai, 458 - Centro, Itajaí/SC"
            break

    for r in paras[4].runs[1:]:
        r.text = ""
    paras[4].runs[0].text = "{" + ", ".join(a[1] for a in autores) + "}@edu.univali.br"

    def set_abstract(p, text):
        for r in p.runs[2:]:
            r.text = ""
        p.runs[2].text = text
    set_abstract(paras[5], text_of("Abstract"))
    set_abstract(paras[6], text_of("Resumo"))

    # --- corpo, por título do template
    heads = [(i, p) for i, p in enumerate(paras) if p.style.name.startswith("Heading")]
    used, video_seen = set(), False
    for idx, (i, h) in enumerate(heads):
        nxt = heads[idx + 1][0] if idx + 1 < len(heads) else len(paras)
        title, lvl = h.text.strip(), h.style.name
        old = paras[i + 1:nxt]
        if lvl == "Heading 4":  # exemplos do template
            h._p.getparent().remove(h._p)
            for bp in old:
                bp._p.getparent().remove(bp._p)
            continue
        if title == "Vídeo de Apresentação" and lvl == "Heading 1":
            video_seen = True
        key = title
        if video_seen and lvl == "Heading 2":
            key = "Vídeo: " + title
        if title in TITLE_ONLY or (video_seen and title == "Etapa 2"):
            new = []
        elif key in sections:
            new = elements(sections[key]); used.add(key)
        else:
            continue
        for bp in old:
            bp._p.getparent().remove(bp._p)
        cur = h._p
        for el in new:
            cur.addnext(el); cur = el

    ignorados = set(sections) - used - {"Título", "Resumo", "Abstract"}
    if ignorados:
        print("AVISO: seções sem destino no template:", ", ".join(sorted(ignorados)), file=sys.stderr)

    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    doc.save(out)
    print("salvo:", out)


def read_autores(path):
    autores = []
    for line in open(path, encoding="utf-8"):
        line = line.strip()
        if line and not line.startswith("#") and "|" in line:
            nome, user = [x.strip() for x in line.split("|", 1)]
            autores.append((nome, user))
    if len(autores) < 1:
        sys.exit("autores vazio: " + path)
    return autores


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--out", default=os.path.join(HERE, "build", "relatorio-etapa1.docx"))
    ap.add_argument("--sections", default=os.path.join(HERE, "etapa1"))
    ap.add_argument("--autores", default=os.path.join(HERE, "autores.local.md"))
    a = ap.parse_args()
    if not os.path.exists(a.autores):
        sys.exit("Crie relatorio/autores.local.md a partir de autores.exemplo.md")
    build(a.template, a.out, parse_sections(a.sections), read_autores(a.autores))
